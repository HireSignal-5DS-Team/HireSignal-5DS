r"""Labelling kit for 40 real Djinni CV x job pairs (T03 Part C). No model is used anywhere in this script.

Reads data/raw/djinni/ (parquet) and data/labels/requirements_draft.csv (hand-written DRAFT). Writes:
  data/processed/jobs/J1..J4.txt      job text with its Djinni id
  data/processed/cvs/<cv_id>.txt      CV text (Djinni id), Position / Experience / English header + CV column
  data/labels/pairs.csv               40 pairs with the selection stratum (NOT shown to labelers)
  data/labels/labelling_sheet_A.xlsx, _B.xlsx  blind sheets, one per labeler: one row per (pair, requirement)

CV selection uses only Primary Keyword and Experience Years, seeded random inside each stratum:
  strong  3  main keyword, experience >= job years
  partial 4  main keyword, 1 <= experience < job years
  weak    3  2 x main keyword with experience <= 0.5 + 1 x adjacent keyword (any experience)
CVs with a rough personal-data hit (same regexes as scripts/djinni_stats.py) or outside 400-3000 chars are skipped.

Usage: .venv\\Scripts\\python scripts\\build_labelling_kit.py
"""
import csv
import random
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from djinni_stats import PII, has_phone  # noqa: E402

RAW = ROOT / "data/raw/djinni"
SEED = 20261018
LABELERS = ("A", "B")
# job, Djinni job id, main CV keyword, adjacent keyword (weak stratum), job years (Djinni "Exp Years" field)
JOBS = [("J1", "617ef5c1-7b6c-5ef6-b0bb-9937e5879294", "Data Science", "Data Analyst", 5),
        ("J2", "1d864d06-e9ee-55a8-a42e-3ff4da68d7a8", "Data Engineer", "SQL", 3),
        ("J3", "f87c22c3-525f-5999-8a64-63cf49ae7a90", "Data Science", "Python", 3),
        ("J4", "0dc63314-bb5e-551c-a012-65d5850b4938", "Python", "Data Engineer", 3)]
PII_RE = re.compile("|".join(f"(?:{p})" for p in PII.values()).replace("(?i)", ""), re.I)


def clean(s):
    return "" if s is None or (isinstance(s, float) and pd.isna(s)) else str(s).replace("\r\n", "\n").strip()


def pick(df, rng, n, taken):
    ids = sorted(set(df["id"]) - taken)
    out = rng.sample(ids, n)
    taken.update(out)
    return out


def write_sheet(order, reqs, who):
    """data/labels/labelling_sheet_<who>.xlsx: only this labeler's columns. Returns the number of requirement rows."""
    wb = Workbook()
    ws = wb.active
    ws.title = "requirements"
    ws.append(["pair_id", "job", "cv_file", "req_id", "kind", "category", "min_years", "requirement",
               f"{who}_decision", f"{who}_cv_quote", f"{who}_notes"])
    for p in order:
        for q in (q for q in reqs if q["job"] == p["job"]):
            ws.append([p["pair_id"], p["job"], f"cvs/{p['cv_id']}.txt", q["req_id"], q["kind"], q["category"],
                       q["min_years"], q["requirement"], "", "", ""])
    dv = DataValidation(type="list", formula1='"MET,MISSING,UNCERTAIN"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"I2:I{ws.max_row}")

    ws2 = wb.create_sheet("pairs")
    ws2.append(["pair_id", "job", "cv_file", f"{who}_outcome", f"{who}_reason"])
    for p in order:
        ws2.append([p["pair_id"], p["job"], f"cvs/{p['cv_id']}.txt", "", ""])
    dv2 = DataValidation(type="list", formula1='"invite,review,reject"', allow_blank=True)
    ws2.add_data_validation(dv2)
    dv2.add(f"D2:D{ws2.max_row}")

    ws3 = wb.create_sheet("job_requirements")
    ws3.append(["status", "job", "req_id", "kind", "category", "min_years", "requirement", "source_sentence"])
    for q in reqs:
        ws3.append([q[k] for k in ("status", "job", "req_id", "kind", "category", "min_years", "requirement", "source_sentence")])

    widths = {"requirement": 60, f"{who}_cv_quote": 45, f"{who}_notes": 30, "cv_file": 46, "source_sentence": 70,
              f"{who}_reason": 40}
    for sheet in (ws, ws2, ws3):
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDDDDD")
            sheet.column_dimensions[cell.column_letter].width = widths.get(cell.value, 14)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(ROOT / f"data/labels/labelling_sheet_{who}.xlsx")
    return ws.max_row - 1


def main():
    jobs = pd.read_parquet(next((RAW / "job-descriptions-english").rglob("*.parquet"))).set_index("id")
    cvs = pd.read_parquet(next((RAW / "candidate-profiles-english").rglob("*.parquet")))
    text = cvs["CV"].fillna("")
    ok = text.str.len().between(400, 3000) & ~text.str.contains(PII_RE) & ~text.map(has_phone)
    pool = cvs[ok]
    exp = pool["Experience Years"]

    reqs = list(csv.DictReader(open(ROOT / "data/labels/requirements_draft.csv", encoding="utf-8")))
    rng, taken, pairs = random.Random(SEED), set(), []
    for job, jid, kw, adj, years in JOBS:
        r = jobs.loc[jid]
        desc = clean(r["Long Description"])
        for q in (q for q in reqs if q["job"] == job):  # self-check: every source sentence is verbatim in the job text
            assert q["source_sentence"] in desc, (job, q["req_id"])
        (ROOT / f"data/processed/jobs/{job}.txt").write_text(
            f"Djinni job id: {jid}\nPosition: {r['Position']}\nExperience (Djinni field): {r['Exp Years']}\n"
            f"English level (Djinni field): {r['English Level']}\n\n{desc}\n", encoding="utf-8")
        main_kw = pool["Primary Keyword"] == kw
        strata = [("strong", pick(pool[main_kw & (exp >= years)], rng, 3, taken)),
                  ("partial", pick(pool[main_kw & (exp >= 1) & (exp < years)], rng, 4, taken)),
                  ("weak", pick(pool[main_kw & (exp <= 0.5)], rng, 2, taken)
                   + pick(pool[pool["Primary Keyword"] == adj], rng, 1, taken))]
        for stratum, ids in strata:
            for cid in ids:
                pairs.append({"pair_id": f"P{len(pairs) + 1:02d}", "job": job, "djinni_job_id": jid, "cv_id": cid,
                              "primary_keyword": pool.loc[pool["id"] == cid, "Primary Keyword"].item(),
                              "experience_years": pool.loc[pool["id"] == cid, "Experience Years"].item(),
                              "selection_stratum": stratum})
    assert len(pairs) == 40 and len({p["cv_id"] for p in pairs}) == 40

    by_id = pool.set_index("id")
    for p in pairs:
        c = by_id.loc[p["cv_id"]]
        (ROOT / f"data/processed/cvs/{p['cv_id']}.txt").write_text(
            f"Position: {clean(c['Position'])}\nExperience years: {c['Experience Years']:g}\n"
            f"English level: {clean(c['English Level'])}\n\n{clean(c['CV'])}\n", encoding="utf-8")
    with open(ROOT / "data/labels/pairs.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(pairs[0]))
        w.writeheader()
        w.writerows(pairs)

    # Blind sheets, one file per labeler: pair order shuffled so the stratum cannot be guessed from the row position
    order = sorted(pairs, key=lambda p: (p["job"], random.Random(SEED + int(p["pair_id"][1:])).random()))
    rows = [write_sheet(order, reqs, who) for who in LABELERS]
    (ROOT / "data/labels/labelling_sheet.xlsx").unlink(missing_ok=True)  # T03 combined sheet (unused), replaced
    print(f"pairs {len(pairs)}; sheet rows {rows[0]} per labeler; strata "
          + str(pd.Series([p["selection_stratum"] for p in pairs]).value_counts().to_dict()))


if __name__ == "__main__":
    main()
