"""Build fake CV files on the fly, so the tests need no real personal data."""

import io
import sys
from pathlib import Path

import docx
import pymupdf
import pytest
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

W, H = A4


def _single_column_pdf() -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    y = H - 60

    def line(text, x=50, size=11, bold=False, right=None):
        nonlocal y
        c.setFont("Helvetica-Bold" if bold else "Helvetica", size)
        c.drawString(x, y, text)
        if right:  # right-aligned date on the same row
            c.drawRightString(W - 50, y, right)
        y -= size + 8

    line("Sami Trabelsi", size=18, bold=True)
    line("sami.trabelsi@example.com | +216 20 123 456 | Tunis")
    y -= 10
    line("PROFESSIONAL EXPERIENCE", bold=True, size=13)
    line("Data Engineer - Vermeg", bold=True, right="2022 - 2024")
    line("Built ETL pipelines in Python and Airflow for banking data.")
    line("Machine Learning Intern - Sofrecom", bold=True, right="2021")
    line("Trained a churn model with scikit-learn on 2M customer records.")
    y -= 10
    line("EDUCATION", bold=True, size=13)
    line("Engineering degree in Data Science - ESPRIT", right="2019 - 2024")
    y -= 10
    line("SKILLS", bold=True, size=13)
    line("Python, SQL, PyTorch, Docker, AWS, Airflow")
    y -= 10
    line("LANGUAGES", bold=True, size=13)
    line("Arabic (native), French (fluent), English (fluent)")
    c.save()
    return buf.getvalue()


def _two_column_pdf() -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setFont("Helvetica-Bold", 20)
    c.drawString(40, H - 50, "Ines Ben Salah")  # full-width banner

    def column(x, y, items):
        for text, bold in items:
            c.setFont("Helvetica-Bold" if bold else "Helvetica", 12 if bold else 10)
            c.drawString(x, y, text)
            y -= 30 if bold else 16

    # Left sidebar (narrow, starts at x=40)
    column(40, H - 110, [
        ("SKILLS", True),
        ("Python", False), ("TensorFlow", False), ("Kubernetes", False), ("PostgreSQL", False),
        ("LANGUAGES", True),
        ("French", False), ("English", False),
        ("CONTACT", False), ("ines.bs@example.com", False),
    ])
    # Main column (wide, starts at x=220) - rows NOT aligned with the sidebar
    column(220, H - 118, [
        ("EXPERIENCE", True),
        ("Backend Developer at Talan Tunisie, 2020 to 2023.", False),
        ("Designed REST APIs serving 50k daily users with FastAPI.", False),
        ("Migrated monolith services to Kubernetes on Azure.", False),
        ("EDUCATION", True),
        ("Master in Software Engineering, INSAT, 2020.", False),
        ("Thesis on distributed caching for microservices.", False),
    ])
    c.save()
    return buf.getvalue()


def _docx() -> bytes:
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "Youssef Hamdi - youssef.h@example.com"
    d.add_heading("Expérience professionnelle", level=1)
    d.add_paragraph("Ingénieur DevOps chez Ooredoo (2021 - 2024)")
    d.add_heading("Compétences :", level=1)
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Cloud"
    table.cell(0, 1).text = "AWS, GCP"
    table.cell(1, 0).text = "Langages"
    table.cell(1, 1).text = "Python, Go"
    d.add_heading("Formation", level=1)
    d.add_paragraph("Diplôme d'ingénieur, ENIT, 2021")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pdf_to_png(pdf_bytes: bytes, dpi: int = 200) -> bytes:
    page = pymupdf.open(stream=pdf_bytes, filetype="pdf")[0]
    return page.get_pixmap(dpi=dpi).tobytes("png")


def _png_to_scanned_pdf(png: bytes) -> bytes:
    """A PDF that only contains a picture of the page: no text layer."""
    pdf = pymupdf.open()
    page = pdf.new_page(width=W, height=H)
    page.insert_image(page.rect, stream=png)
    return pdf.tobytes()


@pytest.fixture(scope="session")
def text_pdf():
    return _single_column_pdf()


@pytest.fixture(scope="session")
def two_column_pdf():
    return _two_column_pdf()


@pytest.fixture(scope="session")
def word_docx():
    return _docx()


@pytest.fixture(scope="session")
def photo_png(text_pdf):
    return _pdf_to_png(text_pdf)


@pytest.fixture(scope="session")
def scanned_pdf(photo_png):
    return _png_to_scanned_pdf(photo_png)


@pytest.fixture(scope="session")
def noisy_photo(photo_png):
    """The clean photo with salt-and-pepper noise, like a bad scan (real case: samples/1.png)."""
    import random
    from PIL import Image
    random.seed(0)
    img = Image.open(io.BytesIO(photo_png)).convert("L")
    px = img.load()
    for _ in range(int(img.width * img.height * 0.08)):
        x, y = random.randrange(img.width), random.randrange(img.height)
        px[x, y] = random.choice((0, 90, 160))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


@pytest.fixture(scope="session")
def two_column_photo(two_column_pdf):
    return _pdf_to_png(two_column_pdf)
