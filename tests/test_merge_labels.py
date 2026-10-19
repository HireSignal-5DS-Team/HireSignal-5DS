"""scripts/merge_labels.py on filled copies of the real blind sheets (tmp files, the originals stay empty)."""
import shutil
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from merge_labels import agreement, merge  # noqa: E402


def test_kappa_hand_computed():
    # 4 rows: agree on 3. p_o = 0.75; A: MET 2, MISSING 2; B: MET 1, MISSING 3 -> p_e = (2*1 + 2*3)/16 = 0.5
    out = agreement(["MET", "MET", "MISSING", "MISSING"], ["MET", "MISSING", "MISSING", "MISSING"],
                    {"MET", "MISSING", "UNCERTAIN"})
    assert out["percent_agreement"] == 0.75 and out["cohen_kappa"] == 0.5


def test_unlabelled_rows_are_not_scored_and_one_class_has_no_kappa():
    out = agreement(["MET", "", "MET"], ["MET", "MET", ""], {"MET", "MISSING", "UNCERTAIN"})
    assert out["n"] == 1 and out["percent_agreement"] == 1.0 and out["cohen_kappa"] is None


@pytest.mark.skipif(not (ROOT / "data/labels/labelling_sheet_A.xlsx").exists(), reason="sheets not built")
def test_merge_real_sheets(tmp_path):
    paths = {}
    for who in ("A", "B"):
        paths[who] = tmp_path / f"{who}.xlsx"
        shutil.copy(ROOT / f"data/labels/labelling_sheet_{who}.xlsx", paths[who])
    first_cv = (ROOT / "data/processed" / load_workbook(paths["A"])["requirements"]["C2"].value).read_text(encoding="utf-8")
    quote = first_cv.splitlines()[0]  # "Position: ..." is in the CV file
    for who, (d2, d3, out2) in {"A": ("MET", "MISSING", "reject"), "B": ("MET", "UNCERTAIN", "review")}.items():
        wb = load_workbook(paths[who])
        ws, wp = wb["requirements"], wb["pairs"]
        ws["I2"], ws["J2"], ws["I3"] = d2, quote, d3
        wp["D2"] = out2
        if who == "B":
            ws["I4"], ws["J4"] = "MET", "not in the cv at all"
        wb.save(paths[who])
    req, pairs, summary = merge(paths["A"], paths["B"], ROOT / "data/processed/cvs")
    assert len(req) == 280 and len(pairs) == 40
    assert summary["requirements"]["n"] == 2 and summary["requirements"]["percent_agreement"] == 0.5
    assert summary["outcomes"]["n"] == 1 and summary["outcomes"]["percent_agreement"] == 0.0
    assert summary["labelled"]["B"]["requirements"] == 3
    assert any("not an exact substring" in p for p in summary["problems"])
